"""Export tests: service-level builders + the /api/export endpoints."""
import csv
import io
from datetime import date, time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.services.export_service import (
    CSV_COLUMNS,
    build_attendance_csv,
    build_attendance_xlsx,
)


def _rows():
    return [
        {
            "date": date(2026, 10, 2),
            "class_name": "Grade 5",
            "section": "A",
            "subject": "Mathematics",
            "student_id": "STU001",
            "roll_number": "1",
            "name": "Aarav Sharma",
            "status": "present",
            "manually_confirmed": True,
        },
        {
            "date": date(2026, 10, 2),
            "class_name": "Grade 5",
            "section": "A",
            "subject": "Mathematics",
            "student_id": "STU002",
            "roll_number": "2",
            "name": "Diya Patel",
            "status": "needs_review",
            "manually_confirmed": False,
        },
    ]


def _meta():
    return {
        "class_name": "Grade 5",
        "section": "A",
        "subject": "Mathematics",
        "date": date(2026, 10, 2),
        "session_time": time(9, 0),
    }


def _summary():
    return {
        "total": 2, "present": 1, "absent": 0, "excused": 0,
        "needs_review": 1, "attendance_percentage": 50.0,
    }


# --- service level ---------------------------------------------------------
def test_csv_columns_contract():
    assert CSV_COLUMNS == [
        "Date", "Class", "Section", "Subject", "Student ID", "Roll Number",
        "Student Name", "Status", "Manually Confirmed",
    ]


def test_build_csv_headers_and_rows():
    text = build_attendance_csv(_rows())
    parsed = list(csv.DictReader(io.StringIO(text)))
    assert len(parsed) == 2
    assert list(parsed[0].keys()) == CSV_COLUMNS
    assert parsed[0]["Student Name"] == "Aarav Sharma"
    assert parsed[0]["Status"] == "Present"          # human-readable label
    assert parsed[1]["Status"] == "Needs Review"
    assert parsed[0]["Manually Confirmed"] == "Yes"
    assert parsed[1]["Manually Confirmed"] == "No"


def test_build_xlsx_valid_zip_bytes():
    data = build_attendance_xlsx(_meta(), _rows(), _summary())
    assert isinstance(data, bytes)
    assert data[:2] == b"PK"  # zip container

    wb = load_workbook(io.BytesIO(data))
    assert "Attendance" in wb.sheetnames
    ws = wb["Attendance"]
    # Find the table header row (first two cells "Date", "Class" — the meta
    # section above also has a "Date" label row, so match on both cells).
    header = None
    for row in ws.iter_rows(values_only=True):
        if row[0] == "Date" and row[1] == "Class":
            header = list(row)
            break
    assert header == CSV_COLUMNS, f"header row mismatch: {header}"


# --- API level -------------------------------------------------------------
@pytest.fixture()
def finalized_session_id(client: TestClient, teacher_headers: dict,
                         test_class: dict) -> int:
    r = client.post(
        "/api/students",
        json={"class_id": test_class["id"], "student_id": "X1",
              "roll_number": "1", "name": "Export Student"},
        headers=teacher_headers,
    )
    assert r.status_code == 201, r.text
    student_id = r.json()["id"]

    r = client.post(
        "/api/attendance/sessions",
        json={"class_id": test_class["id"], "date": "2026-10-02",
              "session_time": "09:00:00", "subject": "Mathematics"},
        headers=teacher_headers,
    )
    assert r.status_code in (200, 201), r.text
    sid = r.json()["id"]

    r = client.put(
        f"/api/attendance/sessions/{sid}/records",
        json={"records": [{"student_id": student_id, "status": "present",
                           "manually_confirmed": True}]},
        headers=teacher_headers,
    )
    assert r.status_code == 200, r.text
    r = client.post(f"/api/attendance/sessions/{sid}/finalize", json={},
                    headers=teacher_headers)
    assert r.status_code == 200, r.text
    return sid


def test_export_csv_endpoint(client: TestClient, teacher_headers: dict,
                             finalized_session_id: int):
    r = client.get(
        "/api/export/attendance",
        params={"session_id": finalized_session_id, "format": "csv"},
        headers=teacher_headers,
    )
    assert r.status_code == 200, f"export csv failed: {r.status_code} {r.text}"
    assert "text/csv" in r.headers["content-type"]
    parsed = list(csv.DictReader(io.StringIO(r.text)))
    assert list(parsed[0].keys()) == CSV_COLUMNS
    assert parsed[0]["Student Name"] == "Export Student"
    assert parsed[0]["Status"] == "Present"


def test_export_xlsx_endpoint(client: TestClient, teacher_headers: dict,
                              finalized_session_id: int):
    r = client.get(
        "/api/export/attendance",
        params={"session_id": finalized_session_id, "format": "xlsx"},
        headers=teacher_headers,
    )
    assert r.status_code == 200, f"export xlsx failed: {r.status_code} {r.text}"
    assert r.content[:2] == b"PK"
    wb = load_workbook(io.BytesIO(r.content))
    assert "Attendance" in wb.sheetnames


def test_export_requires_auth(client: TestClient, finalized_session_id: int):
    r = client.get("/api/export/attendance",
                   params={"session_id": finalized_session_id, "format": "csv"})
    assert r.status_code == 401
