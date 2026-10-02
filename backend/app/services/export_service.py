"""CSV / XLSX export builders for attendance data."""
import csv
import io
from datetime import date, time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

CSV_COLUMNS = [
    "Date",
    "Class",
    "Section",
    "Subject",
    "Student ID",
    "Roll Number",
    "Student Name",
    "Status",
    "Manually Confirmed",
]


def status_label(status: str) -> str:
    return status.replace("_", " ").title()


def build_attendance_csv(rows: list[dict]) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=CSV_COLUMNS)
    writer.writeheader()
    for r in rows:
        writer.writerow(
            {
                "Date": r["date"].isoformat()
                if isinstance(r["date"], date)
                else r["date"],
                "Class": r["class_name"],
                "Section": r["section"],
                "Subject": r["subject"],
                "Student ID": r["student_id"],
                "Roll Number": r["roll_number"],
                "Student Name": r["name"],
                "Status": status_label(r["status"]),
                "Manually Confirmed": "Yes" if r["manually_confirmed"] else "No",
            }
        )
    return buf.getvalue()


def build_attendance_xlsx(
    meta: dict, rows: list[dict], summary: dict
) -> bytes:
    """Formatted workbook: title, class info, summary, then the record table."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    title_font = Font(size=16, bold=True, color="1F4E5F")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E5F")
    label_font = Font(bold=True)
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:I1")
    ws["A1"] = "Classroom Attendance Report"
    ws["A1"].font = title_font

    info = [
        ("Class", meta.get("class_name", "")),
        ("Section", meta.get("section", "")),
        ("Subject", meta.get("subject", "")),
        (
            "Date",
            meta["date"].isoformat()
            if isinstance(meta.get("date"), date)
            else meta.get("date", ""),
        ),
        (
            "Session Time",
            meta["session_time"].strftime("%H:%M")
            if isinstance(meta.get("session_time"), time)
            else meta.get("session_time", ""),
        ),
    ]
    row = 3
    for label, value in info:
        ws.cell(row=row, column=1, value=label).font = label_font
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Summary").font = Font(size=13, bold=True)
    row += 1
    for key in ("total", "present", "absent", "excused", "needs_review"):
        ws.cell(row=row, column=1, value=key.replace("_", " ").title()).font = label_font
        ws.cell(row=row, column=2, value=summary.get(key, 0))
        row += 1
    ws.cell(row=row, column=1, value="Attendance %").font = label_font
    ws.cell(row=row, column=2, value=summary.get("attendance_percentage", 0.0))

    row += 2
    table_start = row
    for col, name in enumerate(CSV_COLUMNS, start=1):
        cell = ws.cell(row=row, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = border
    row += 1
    for r in rows:
        values = [
            r["date"].isoformat() if isinstance(r["date"], date) else r["date"],
            r["class_name"],
            r["section"],
            r["subject"],
            r["student_id"],
            r["roll_number"],
            r["name"],
            status_label(r["status"]),
            "Yes" if r["manually_confirmed"] else "No",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = border
        row += 1

    widths = [14, 22, 10, 18, 14, 12, 26, 14, 18]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=table_start, column=col).column_letter].width = w
    ws.freeze_panes = ws.cell(row=table_start + 1, column=1).coordinate

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
